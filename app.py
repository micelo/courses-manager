import os
import sqlite3
from datetime import datetime, date
from pathlib import Path
from typing import Dict, Optional, Tuple, List

from flask import (
    Flask,
    flash,
    g,
    redirect,
    render_template,
    request,
    session,
    url_for,
    current_app,
)
from werkzeug.security import check_password_hash, generate_password_hash

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover - handled at runtime
    raise RuntimeError(
        "openpyxl is required to import registrations from Excel files"
    ) from exc


def create_app() -> Flask:
    app = Flask(__name__)
    base_dir = Path(__file__).resolve().parent
    app.config["SECRET_KEY"] = os.environ.get("COURSES_MANAGER_SECRET_KEY",
                                              "bb6510346d1d2d1fb6aa5802c04e625b05c4e29141ab1320b86705dff3cf2874")
    app.config["DATABASE"] = str(base_dir / "app.db")
    app.config["EXCEL_FILE"] = str(base_dir / "registrations.xlsx")

    with app.app_context():
        _initialize_database()
        _ensure_default_user()

    register_routes(app)
    return app


def get_db() -> sqlite3.Connection:
    if "db" not in g:
        g.db = sqlite3.connect(
            current_app.config["DATABASE"],
            detect_types=sqlite3.PARSE_DECLTYPES,
        )
        g.db.row_factory = sqlite3.Row
    return g.db


def close_db(_: Optional[BaseException] = None) -> None:
    db = g.pop("db", None)
    if db is not None:
        db.close()


def _initialize_database() -> None:
    db = sqlite3.connect(current_app.config["DATABASE"])
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE NOT NULL,
            password_hash TEXT NOT NULL
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS registrations (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            surname TEXT NOT NULL,
            registration_date TEXT,
            id_number TEXT NOT NULL,
            hrda_number TEXT,
            course_name TEXT NOT NULL,
            employment_category TEXT,
            email TEXT,
            phone TEXT,
            status TEXT NOT NULL DEFAULT 'PENDING',
            UNIQUE(id_number, course_name)
)
        """
    )
    db.commit()
    db.close()


def _ensure_default_user() -> None:
    username = os.environ.get("COURSES_MANAGER_DEFAULT_USERNAME", "admin")
    password = os.environ.get("COURSES_MANAGER_DEFAULT_PASSWORD", "changeme")

    db = sqlite3.connect(current_app.config["DATABASE"])
    db.row_factory = sqlite3.Row
    user = db.execute("SELECT id FROM users WHERE username = ?", (username,)).fetchone()
    if user is None:
        db.execute(
            "INSERT INTO users (username, password_hash) VALUES (?, ?)",
            (username, generate_password_hash(password)),
        )
        db.commit()
    db.close()


COURSES: List[str] = [
    "Αναδυόμενες Τεχνολογίες και Δεξιότητες για Ανάπτυξη πιο Έξυπνων και Πράσινων Πόλεων",
    "Αναδυόμενες Τεχνολογίες και Δεξιότητες για Ανάπτυξη πιο Έξυπνων και Πράσινων Πόλεων (Μόνο για Ανέργους)",
    "Αντλίες Θερμότητας στην Πράσινη Μετάβαση",
    "Αντλίες Θερμότητας στην Πράσινη Μετάβαση (Μόνο για Ανέργους)",
    "Αποθήκευση Ενέργειας: Πολύπλευρος Ρόλος στο Σύγχρονο Ηλεκτρικό Δίκτυο",
    "Αποθήκευση Ενέργειας: Πολύπλευρος Ρόλος στο Σύγχρονο Ηλεκτρικό Δίκτυο (Μόνο για Ανέργους)",
    "Εισαγωγή στα Κτήρια με Σχεδόν Μηδενική Κατανάλωση Ενέργειας",
    "Εισαγωγή στα Κτήρια με Σχεδόν Μηδενική Κατανάλωση Ενέργειας (Μόνο για Ανέργους)",
    "Έλεγχος και Επιθεώρηση Φωτοβολταϊκών Συστημάτων",
    "Έλεγχος και Επιθεώρηση Φωτοβολταϊκών Συστημάτων (Μόνο για Ανέργους)",
    "Έννοιες, Πλαίσιο και Πολιτικές για μια Πράσινη Οικονομία χωρίς Αποκλεισμούς",
    "Ενσωμάτωση της Ηλεκτρικής Κινητικότητας στο Σύστημα Ηλεκτρισμού",
    "Ενσωμάτωση της Ηλεκτρικής Κινητικότητας στο Σύστημα Ηλεκτρισμού (Μόνο για Ανέργους)",
    "Έξυπνα Ηλεκτρικά Δίκτυα: Τεχνολογίες, Διαχείριση και Πρακτικές Εφαρμογές",
    "Έξυπνα Ηλεκτρικά Δίκτυα: Τεχνολογίες, Διαχείριση και Πρακτικές Εφαρμογές (Μόνο για Ανέργους)",
    "Έξυπνα Συστήματα Διαχείρισης Ενέργειας",
    "Έξυπνα Συστήματα Διαχείρισης Ενέργειας (Μόνο για Ανέργους)",
    "Κανόνες Αγοράς Ηλεκτρισμού",
    "Κυκλική Οικονομία ΦΒ Συστημάτων: Ευκαιρίες για Επισκευή, Επαναχρησιμοποίηση ή Ανακύκλωση",
    "Κυκλική Οικονομία ΦΒ Συστημάτων: Ευκαιρίες για Επισκευή, Επαναχρησιμοποίηση ή Ανακύκλωση (Μόνο για Ανέργους)",
]

HEADER_ALIASES = {
    "name": (
        "name", "first name", "όνομα",
        "usr_first_name"
    ),
    "surname": (
        "surname", "last name", "επώνυμο", "επίθετο",
        "usr_last_name"
    ),
    "registration_date": (
        "date", "registration date", "ημερομηνία",
        "applic_creation_date"
    ),
    "id_number": (
        "id", "id number", "identification number", "αρ. ταυτότητας", "ταυτότητα", "civil id"
    ),
    "hrda_number": (
        "hrda no.", "hrda number", "αρ. αναδ", "αρ. ανaδ", "αρ. ανaΔ", "hrda no"
    ),
    "course_name": (
        "name of the course", "course name", "program name", "όνομα προγράμματος", "πρόγραμμα",
        "saa_title",
        "SAA_TITLE"
    ),
    "employment_category": (
        "employment category", "κατηγορία απασχόλησης",
        "saa_employment_status"
    ),
    "email": (
        "email", "e-mail", "usr_email"
    ),
    "phone": (
        "phone", "phone number", "telephone", "τηλέφωνο", "usr_tel"
    ),
}

REQUIRED_COLUMNS = {"name", "surname", "id_number", "course_name"}


def register_routes(app: Flask) -> None:
    @app.before_request
    def load_db() -> None:
        get_db()

    app.teardown_appcontext(close_db)

    def login_required(view):
        from functools import wraps

        @wraps(view)
        def wrapped_view(**kwargs):
            if session.get("user_id") is None:
                flash("Please log in to access the dashboard.", "warning")
                return redirect(url_for("login"))
            return view(**kwargs)

        return wrapped_view

    @app.route("/login", methods=["GET", "POST"])
    def login():
        if request.method == "POST":
            username = request.form.get("username", "").strip()
            password = request.form.get("password", "")

            db = get_db()
            user = db.execute("SELECT * FROM users WHERE username = ?", (username,)).fetchone()
            if user and check_password_hash(user["password_hash"], password):
                session.clear()
                session["user_id"] = user["id"]
                session["username"] = user["username"]
                try:
                    imported = sync_from_excel()
                    if imported:
                        flash(f"Imported {imported} registrations from Excel.", "success")
                except Exception as exc:  # pragma: no cover - defensive logging
                    current_app.logger.exception("Failed to import registrations: %s", exc)
                    flash("Login succeeded but importing the Excel file failed. Check the logs for details.", "danger")
                return redirect(url_for("dashboard"))

            flash("Invalid username or password.", "danger")
        return render_template("login.html")

    @app.route("/logout")
    def logout():
        session.clear()
        flash("You have been logged out.", "info")
        return redirect(url_for("login"))

    def _parse_filters() -> Tuple[int, str, str]:
        try:
            page = int(request.args.get("page", 1))
        except ValueError:
            page = 1
        page = max(page, 1)

        selected_course = request.args.get("course", "").strip()
        selected_date = request.args.get("date", "").strip()

        return page, selected_course, selected_date

    def _fetch_registrations(
        base_condition: str,
        base_params: Tuple[object, ...],
        selected_course: str,
        selected_date: str,
        page: int,
    ) -> Tuple[List[sqlite3.Row], bool]:
        per_page = 20
        offset = (page - 1) * per_page

        conditions = [base_condition]
        params: List[object] = list(base_params)

        if selected_course:
            conditions.append("course_name = ?")
            params.append(selected_course)

        if selected_date:
            conditions.append("registration_date = ?")
            params.append(selected_date)

        where_clause = " WHERE " + " AND ".join(conditions) if conditions else ""

        query = f"""
            SELECT id, name, surname, registration_date, id_number, hrda_number,
                   course_name, employment_category, email, phone, status
            FROM registrations
            {where_clause}
            ORDER BY COALESCE(registration_date, '') ASC, id ASC
            LIMIT ? OFFSET ?
        """

        params.extend([per_page, offset])

        db = get_db()
        registrations = db.execute(query, tuple(params)).fetchall()

        # Determine whether there might be more records for pagination
        has_next = len(registrations) == per_page
        return registrations, has_next

    @app.route("/")
    @login_required
    def dashboard():
        page, selected_course, selected_date = _parse_filters()
        registrations, has_next = _fetch_registrations(
            "status != ?",
            ("CALLED",),
            selected_course,
            selected_date,
            page,
        )

        return render_template(
            "dashboard.html",
            registrations=registrations,
            page=page,
            has_next=has_next,
            selected_course=selected_course,
            selected_date=selected_date,
            courses=COURSES,
            allow_call_action=True,
            pagination_endpoint="dashboard",
            prev_url=url_for(
                "dashboard",
                page=page - 1,
                course=selected_course or None,
                date=selected_date or None,
            ) if page > 1 else None,
            next_url=url_for(
                "dashboard",
                page=page + 1,
                course=selected_course or None,
                date=selected_date or None,
            ) if has_next else None,
        )

    @app.route("/called")
    @login_required
    def called_participants():
        page, selected_course, selected_date = _parse_filters()
        registrations, has_next = _fetch_registrations(
            "status = ?",
            ("CALLED",),
            selected_course,
            selected_date,
            page,
        )

        return render_template(
            "dashboard.html",
            registrations=registrations,
            page=page,
            has_next=has_next,
            selected_course=selected_course,
            selected_date=selected_date,
            courses=COURSES,
            allow_call_action=False,
            pagination_endpoint="called_participants",
            prev_url=url_for(
                "called_participants",
                page=page - 1,
                course=selected_course or None,
                date=selected_date or None,
            ) if page > 1 else None,
            next_url=url_for(
                "called_participants",
                page=page + 1,
                course=selected_course or None,
                date=selected_date or None,
            ) if has_next else None,
        )

    @app.post("/registrations/<int:registration_id>/call")
    @login_required
    def mark_called(registration_id: int):
        db = get_db()
        updated = db.execute(
            "UPDATE registrations SET status = 'CALLED' WHERE id = ?",
            (registration_id,),
        )
        db.commit()
        if updated.rowcount:
            flash("Marked participant as called.", "success")
        else:
            flash("Participant not found or already marked.", "warning")
        return redirect(url_for("dashboard"))


def normalize_header(value: Optional[str]) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip().lower()
    if not text:
        return None
    for canonical, aliases in HEADER_ALIASES.items():
        if text in aliases:
            return canonical
    return None


def normalize_date(value: Optional[object]) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return text


def sync_from_excel() -> int:
    excel_path = Path(current_app.config["EXCEL_FILE"])
    if not excel_path.exists():
        current_app.logger.warning("Excel file '%s' not found. Skipping import.", excel_path)
        return 0

    workbook = load_workbook(excel_path, data_only=True)
    records = {}

    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration:
            continue

        column_map: Dict[str, int] = {}
        for index, header in enumerate(header_row):
            canonical = normalize_header(header)
            if canonical:
                column_map[canonical] = index

        missing = REQUIRED_COLUMNS - column_map.keys()
        if missing:
            current_app.logger.info(
                "Skipping sheet '%s' because required columns %s were not found.",
                sheet.title,
                ", ".join(sorted(missing)),
            )
            continue

        for row in rows:
            if row is None or not any(cell not in (None, "") for cell in row):
                continue

            def get_value(field: str) -> Optional[str]:
                index = column_map.get(field)
                if index is None:
                    return None
                raw = row[index]
                if field == "registration_date":
                    return normalize_date(raw)
                if raw is None:
                    return None
                text = str(raw).strip()
                return text if text else None

            participant_id = get_value("id_number")
            program_name = get_value("course_name")
            if not participant_id or not program_name:
                continue

            key = (participant_id, program_name)
            records[key] = {
                "name": get_value("name") or "",
                "surname": get_value("surname") or "",
                "registration_date": get_value("registration_date"),
                "id_number": participant_id,
                "hrda_number": get_value("hrda_number"),
                "course_name": program_name,
                "employment_category": get_value("employment_category"),
                "email": get_value("email"),
                "phone": get_value("phone"),
            }

    if not records:
        return 0

    db = get_db()
    imported = 0
    for entry in records.values():
        placeholders = (
            entry["name"],
            entry["surname"],
            entry["registration_date"],
            entry["id_number"],
            entry["hrda_number"],
            entry["course_name"],
            entry["employment_category"],
            entry["email"],
            entry["phone"],
        )
        db.execute(
            """
            INSERT INTO registrations (
    name, surname, registration_date, id_number, hrda_number,
    course_name, employment_category, email, phone
) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
ON CONFLICT(id_number, course_name) DO UPDATE SET
    name = excluded.name,
    surname = excluded.surname,
    registration_date = excluded.registration_date,
    hrda_number = excluded.hrda_number,
    employment_category = excluded.employment_category,
    email = excluded.email,
    phone = excluded.phone
            """,
            placeholders,
        )
        imported += 1

    db.commit()
    return imported


app = create_app()

if __name__ == "__main__":
    app.run(port=8000, debug=True)
